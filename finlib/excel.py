"""Build a styled, multi-sheet Excel report (returned as bytes) for download."""

from __future__ import annotations

import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "1F3864"
BLUE = "2E75B6"
GREEN = "548235"
RED = "C00000"
LIGHT = "D9E1F2"
WHITE = "FFFFFF"

_MONEY = "#,##0.00"
_INT = "#,##0"
_DATE = "yyyy-mm-dd"
_thin = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_TABLE_STYLE = TableStyleInfo(
    name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
)


def _style_table_sheet(ws, df: pd.DataFrame, table_name: str, tab_color: str) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab_color
    nrows, ncols = df.shape
    ws.add_table(_table(f"A1:{get_column_letter(ncols)}{nrows + 1}", table_name))
    ws.freeze_panes = "A2"
    for idx, col in enumerate(df.columns, start=1):
        letter = get_column_letter(idx)
        if col == "date":
            fmt = _DATE
        elif pd.api.types.is_float_dtype(df[col].dtype):
            fmt = _MONEY
        elif pd.api.types.is_integer_dtype(df[col].dtype):
            fmt = _INT
        else:
            fmt = None
        if fmt:
            for cell in ws[letter][1:]:
                cell.number_format = fmt
        width = max([len(str(col))] + df[col].astype(str).str.len().tolist())
        ws.column_dimensions[letter].width = min(max(width + 2, 10), 55)


def _table(ref: str, name: str) -> Table:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = _TABLE_STYLE
    return table


def _kpi_card(ws, label_cell, value_cell, title, value, color, fmt=_MONEY) -> None:
    lab = ws[label_cell]
    lab.value = title
    lab.font = Font(bold=True, color=WHITE, size=10)
    lab.fill = PatternFill("solid", fgColor=color)
    lab.alignment = Alignment(horizontal="center", vertical="center")
    lab.border = _BORDER
    val = ws[value_cell]
    val.value = value
    val.font = Font(bold=True, size=14, color=color)
    val.alignment = Alignment(horizontal="center", vertical="center")
    val.number_format = fmt
    val.fill = PatternFill("solid", fgColor=LIGHT)
    val.border = _BORDER


def build_excel_report(
    tx: pd.DataFrame,
    monthly: pd.DataFrame,
    categories: pd.DataFrame,
    kpis: dict,
    *,
    currency: str = "PKR",
    period_label: str = "All data",
) -> bytes:
    """Return an .xlsx report (Summary + Transactions + Monthly + Categories) as bytes."""
    book = Workbook()
    book.calculation.fullCalcOnLoad = True

    # --- Summary ------------------------------------------------------------
    ws = book.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY
    ws.column_dimensions["A"].width = 3
    for letter in "BCDEFGHIJ":
        ws.column_dimensions[letter].width = 15

    ws.merge_cells("B2:H2")
    ws["B2"] = "Financial Summary"
    ws["B2"].font = Font(bold=True, size=20, color=NAVY)
    ws.merge_cells("B3:H3")
    ws["B3"] = f"Period: {period_label}   •   Currency: {currency}"
    ws["B3"].font = Font(italic=True, color="595959")

    _kpi_card(ws, "B5", "B6", "Total Income", round(kpis["income"], 2), GREEN)
    _kpi_card(ws, "D5", "D6", "Total Expense", round(kpis["expense"], 2), RED)
    _kpi_card(ws, "F5", "F6", "Net", round(kpis["net"], 2), NAVY)
    _kpi_card(ws, "H5", "H6", "Transactions", int(kpis["count"]), BLUE, fmt=_INT)
    _kpi_card(ws, "B8", "B9", "Avg Expense", round(kpis["avg_expense"], 2), RED)
    _kpi_card(ws, "D8", "D9", "Savings Rate %", round(kpis["savings_rate"], 1), GREEN, fmt="0.0")
    _kpi_card(ws, "F8", "F9", "Closing Balance", round(kpis["closing_balance"], 2), NAVY)
    _kpi_card(ws, "H8", "H9", "Months", int(kpis["months"]), BLUE, fmt=_INT)

    # --- Transactions -------------------------------------------------------
    ws_tx = book.create_sheet("Transactions")
    tx_out = tx.copy()
    if "date" in tx_out:
        tx_out["date"] = pd.to_datetime(tx_out["date"]).dt.date
    _write_df(ws_tx, tx_out)
    if len(tx_out):
        _style_table_sheet(ws_tx, tx_out, "tblTx", NAVY)

    # --- Monthly (with a bar chart) ----------------------------------------
    ws_m = book.create_sheet("Monthly")
    _write_df(ws_m, monthly)
    if len(monthly):
        _style_table_sheet(ws_m, monthly, "tblMonthly", BLUE)
        cols = list(monthly.columns)
        inc_i, exp_i = cols.index("income") + 1, cols.index("expense") + 1
        chart = BarChart()
        chart.type = "col"
        chart.title = "Income vs Expense by Month"
        chart.height, chart.width = 8, 18
        for c in (inc_i, exp_i):
            chart.add_data(
                Reference(ws_m, min_col=c, min_row=1, max_row=len(monthly) + 1),
                titles_from_data=True,
            )
        chart.set_categories(Reference(ws_m, min_col=1, min_row=2, max_row=len(monthly) + 1))
        ws_m.add_chart(chart, f"{get_column_letter(len(cols) + 2)}2")

    # --- Categories (with a pie chart) -------------------------------------
    ws_c = book.create_sheet("Categories")
    _write_df(ws_c, categories)
    if len(categories):
        _style_table_sheet(ws_c, categories, "tblCat", GREEN)
        pie = PieChart()
        pie.title = "Expense by Category"
        pie.height, pie.width = 9, 13
        pie.add_data(Reference(ws_c, min_col=2, min_row=1, max_row=len(categories) + 1),
                     titles_from_data=True)
        pie.set_categories(Reference(ws_c, min_col=1, min_row=2, max_row=len(categories) + 1))
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True
        ws_c.add_chart(pie, f"{get_column_letter(len(categories.columns) + 2)}2")

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _write_df(ws, df: pd.DataFrame) -> None:
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append(list(row))
